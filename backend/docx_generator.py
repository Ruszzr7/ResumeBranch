"""Create an editable Word version of a structured resume."""

from __future__ import annotations

import base64
import math
import re
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .inline_formatting import parse_inline_bold
from .resume_labels import LABELS
from .resume_data import normalize_resume_data

DEFAULT_FONT_SPEC = {
    "latinFont": "Arial",
    "eastAsiaFont": "Microsoft YaHei",
}

_NATIVE_LIST_MARKER_RE = re.compile(
    r"^\s*([（(]?\d{1,2}[）).、．]|[一二三四五六七八九十]+[、.．])\s*(.*)$",
    re.DOTALL,
)


def _module_list_content(value: object) -> str:
    match = _NATIVE_LIST_MARKER_RE.match(str(value or ""))
    return match.group(2) if match else str(value or "")


def _is_fully_bold(value: object) -> bool:
    segments = [segment for segment in parse_inline_bold(_module_list_content(value).strip()) if segment.text]
    return bool(segments) and all(segment.bold for segment in segments)


def _photo_aspect_ratio(photo_bytes: bytes | None, resume_data: dict) -> float:
    """Return the imported photo's width/height ratio with a legacy fallback."""
    try:
        explicit = float((resume_data.get("basics") or {}).get("photo_aspect_ratio"))
        if 0.2 <= explicit <= 3.0:
            return explicit
    except (TypeError, ValueError):
        pass
    if photo_bytes:
        try:
            from PIL import Image
            with Image.open(BytesIO(photo_bytes)) as image:
                width, height = image.size
            if width and height:
                return max(0.2, min(3.0, width / height))
        except Exception:
            pass
    return 21.0 / 26.0


def _other_field_labels(values: dict, labels: dict) -> dict[str, str]:
    custom = values.get("field_labels") if isinstance(values, dict) else {}
    custom = custom if isinstance(custom, dict) else {}
    return {
        "certificates": str(custom["certificates"] if "certificates" in custom else labels["certificates"]),
        "languages": str(custom["languages"] if "languages" in custom else labels["language"]),
    }


def _other_field_value(label: str, values: list[object], separator: str, colon: str) -> str:
    joined = separator.join(str(value) for value in values)
    return f"{label}{colon}{joined}" if label else joined


def _set_cell_borderless(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _set_cell_margins(cell, top=0, start=0, bottom=0, end=0) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    # Write both legacy and directional names.  The document is intentionally
    # compatible with Word 2010+, whose table inheritance can otherwise retain
    # the template's 108-dxa left/right padding despite start/end overrides.
    for margin, value in (
        ("top", top), ("left", start), ("start", start),
        ("bottom", bottom), ("right", end), ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_fixed_table_widths(table, widths_mm: tuple[float, ...]) -> None:
    """Keep Word from stretching a compact table back to the page edge."""
    table.autofit = False
    table_properties = table._tbl.tblPr

    layout = table_properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(round(sum(widths_mm) * 1440 / 25.4)))

    table_indent = table_properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), "0")

    for column, width in zip(table.columns, widths_mm):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)


def _add_floating_picture(
    paragraph,
    image_bytes: bytes,
    *,
    width_mm: float,
    height_mm: float,
    x_mm: float,
    y_mm: float,
) -> None:
    """Insert a page-positioned picture without putting it in the basics table."""
    inline_shape = paragraph.add_run().add_picture(
        BytesIO(image_bytes),
        width=Mm(width_mm),
        height=Mm(height_mm),
    )
    inline = inline_shape._inline
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    frame_locks = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    if extent is None or doc_pr is None or frame_locks is None or graphic is None:
        return

    inline.tag = qn("wp:anchor")
    for key, value in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", "251658240"),
        ("behindDoc", "0"),
        ("locked", "0"),
        ("layoutInCell", "0"),
        ("allowOverlap", "1"),
    ):
        inline.set(key, value)

    for child in list(inline):
        inline.remove(child)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "page")
    horizontal_offset = OxmlElement("wp:posOffset")
    horizontal_offset.text = str(round(max(0.0, x_mm) * 36000))
    position_h.append(horizontal_offset)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "page")
    vertical_offset = OxmlElement("wp:posOffset")
    vertical_offset.text = str(round(max(0.0, y_mm) * 36000))
    position_v.append(vertical_offset)
    effect_extent = OxmlElement("wp:effectExtent")
    for side in ("l", "t", "r", "b"):
        effect_extent.set(side, "0")
    wrap_none = OxmlElement("wp:wrapNone")

    inline.extend((simple_pos, position_h, position_v, extent, effect_extent, wrap_none, doc_pr, frame_locks, graphic))


def _disable_document_grid(section) -> None:
    """Prevent Word's East Asian document grid from inflating compact lines."""
    section_properties = section._sectPr
    document_grid = section_properties.find(qn("w:docGrid"))
    if document_grid is not None:
        section_properties.remove(document_grid)


def _configure_chinese_document(document) -> None:
    """Make Word use Simplified-Chinese punctuation and theme rules."""
    settings = document.settings._element
    character_spacing = settings.find(qn("w:characterSpacingControl"))
    if character_spacing is None:
        character_spacing = OxmlElement("w:characterSpacingControl")
        settings.insert(0, character_spacing)
    # Keep full-width punctuation at its authored width, matching Chromium's
    # preview/PDF layout instead of letting Word compress punctuation runs.
    character_spacing.set(qn("w:val"), "doNotCompress")

    theme_language = settings.find(qn("w:themeFontLang"))
    if theme_language is None:
        theme_language = OxmlElement("w:themeFontLang")
        settings.append(theme_language)
    theme_language.set(qn("w:val"), "en-US")
    theme_language.set(qn("w:eastAsia"), "zh-CN")

def _set_language(run_properties) -> None:
    language = run_properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        run_properties.append(language)
    language.set(qn("w:val"), "en-US")
    language.set(qn("w:eastAsia"), "zh-CN")


def _set_font(run, size: float, bold: bool = False, color: str = "212529", fonts: dict | None = None) -> None:
    font_spec = fonts or DEFAULT_FONT_SPEC
    latin_font = font_spec["latinFont"]
    east_asia_font = font_spec["eastAsiaFont"]
    run.font.name = latin_font
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{script}"), latin_font)
    r_fonts.set(qn("w:eastAsia"), east_asia_font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run_properties = run._element.get_or_add_rPr()
    _set_language(run_properties)
    spacing = run_properties.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        run_properties.append(spacing)
    spacing.set(qn("w:val"), "0")
    kerning = run_properties.find(qn("w:kern"))
    if kerning is None:
        kerning = OxmlElement("w:kern")
        run_properties.append(kerning)
    kerning.set(qn("w:val"), "0")


def _add_markdown_runs(
    paragraph,
    text: object,
    size: float,
    color: str = "212529",
    fonts: dict | None = None,
    *,
    base_bold: bool = False,
) -> None:
    for segment in parse_inline_bold(str(text or "")):
        _set_font(
            paragraph.add_run(segment.text),
            size,
            bold=base_bold or segment.bold,
            color=color,
            fonts=fonts,
        )


def _paragraph_spacing(paragraph, *, before=0, after=0, line=1.15) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    overflow_punct = p_pr.find(qn("w:overflowPunct"))
    if overflow_punct is None:
        overflow_punct = OxmlElement("w:overflowPunct")
        p_pr.append(overflow_punct)
    # Word enables punctuation overflow when this property is omitted. Keep
    # every line inside the same measured boundary used by HTML/PDF.
    overflow_punct.set(qn("w:val"), "0")
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


WORD_PHOTO_BOTTOM_GAP_MM = 1.5


def _word_estimated_line_count(value: object, font_size_pt: float, width_mm: float) -> int:
    """Estimate Word's wrapped line count from the generated cell width.

    The basics table uses exact paragraph line heights and zero cell margins.
    Measuring the authored text against that same fixed cell width gives the
    DOCX photo boundary a deterministic input without changing the text flow.
    """
    from .layout_config import estimate_text_width_pt

    text = str(value or "").replace("**", "").strip()
    if not text:
        return 0
    available_pt = max(1.0, float(width_mm) * 72.0 / 25.4)
    return max(1, math.ceil(estimate_text_width_pt(text, font_size_pt) / available_pt))


def _word_paragraph_line_height_pt(paragraph, fallback_pt: float) -> float:
    spacing = paragraph._p.get_or_add_pPr().find(qn("w:spacing"))
    if spacing is None:
        return float(fallback_pt)
    line = spacing.get(qn("w:line"))
    if line is None:
        return float(fallback_pt)
    try:
        return float(line) / 20.0
    except (TypeError, ValueError):
        return float(fallback_pt)


def _word_paragraph_spacing_pt(paragraph, key: str) -> float:
    spacing = paragraph._p.get_or_add_pPr().find(qn("w:spacing"))
    if spacing is None:
        return 0.0
    try:
        return float(spacing.get(qn(f"w:{key}")) or 0.0) / 20.0
    except (TypeError, ValueError):
        return 0.0


def _word_bottom_border_extent_pt(paragraph) -> float:
    """Return the Word paragraph's bottom-border gap plus border thickness."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        return 0.0
    bottom = borders.find(qn("w:bottom"))
    if bottom is None or bottom.get(qn("w:val")) in {None, "nil", "none"}:
        return 0.0
    try:
        border_gap = float(bottom.get(qn("w:space")) or 0.0)
    except (TypeError, ValueError):
        border_gap = 0.0
    try:
        # Word border size is expressed in eighths of a point.
        border_width = float(bottom.get(qn("w:sz")) or 0.0) / 8.0
    except (TypeError, ValueError):
        border_width = 0.0
    return border_gap + border_width


def _resolve_word_photo_height_mm(
    desired_height_mm: float,
    *,
    basics_height_pt: float,
    title_paragraph,
) -> float:
    """Cap a DOCX photo before Word's own section divider geometry.

    The page and PDF renderers keep the shared photo height. DOCX has a
    different paragraph/table line box, so use the exact spacing values that
    were written to the Word XML and leave the same physical safety gap before
    the title border. The photo remains page-positioned and cannot move the
    following module.
    """
    title_before_pt = _word_paragraph_spacing_pt(title_paragraph, "before")
    title_line_pt = _word_paragraph_line_height_pt(
        title_paragraph,
        fallback_pt=9.0 * 1.25,
    )
    divider_offset_pt = (
        float(basics_height_pt)
        + title_before_pt
        + title_line_pt
        + _word_bottom_border_extent_pt(title_paragraph)
    )
    max_height_mm = (
        divider_offset_pt * 25.4 / 72.0
        - WORD_PHOTO_BOTTOM_GAP_MM
    )
    return round(max(1.0, min(float(desired_height_mm), max_height_mm)), 2)


def _increase_paragraph_after(paragraph, amount_pt: float) -> None:
    """Add an outer spacing token without discarding inner paragraph spacing."""
    current = paragraph.paragraph_format.space_after
    current_pt = current.pt if current is not None else 0.0
    paragraph.paragraph_format.space_after = Pt(current_pt + amount_pt)


def _date_range(item: dict) -> str:
    values = item.get("date_range") or []
    if isinstance(values, list) and values:
        return " - ".join(str(value) for value in values[:2] if value)
    return " - ".join(value for value in (item.get("start_date", ""), item.get("end_date", "")) if value)


def _set_hanging_indent(paragraph, *, left: float = 5.5, hanging: float = 5.5) -> None:
    """Keep wrapped list lines aligned with the text instead of the marker."""
    paragraph.paragraph_format.left_indent = Mm(left)
    paragraph.paragraph_format.first_line_indent = Mm(-hanging)


WORD_NUMBERED_EXTRA_INDENT_MM = 1.2


def _set_numbered_indent(
    paragraph,
    *,
    text_indent_mm: float,
    marker_gap_mm: float,
    extra_indent_mm: float = 0.0,
) -> None:
    """Right-align the marker and start every wrapped line at one shared stop."""
    text_indent_mm = max(0.0, text_indent_mm) + max(0.0, extra_indent_mm)
    fmt = paragraph.paragraph_format
    fmt.left_indent = Mm(text_indent_mm)
    fmt.first_line_indent = Mm(-text_indent_mm)
    marker_stop = max(0.0, text_indent_mm - marker_gap_mm)
    fmt.tab_stops.add_tab_stop(Mm(marker_stop), WD_TAB_ALIGNMENT.RIGHT)
    fmt.tab_stops.add_tab_stop(Mm(text_indent_mm), WD_TAB_ALIGNMENT.LEFT)


def _set_flush_indent(paragraph) -> None:
    """Match HTML native-marker items: every wrapped line starts at the section edge."""
    paragraph.paragraph_format.left_indent = Mm(0)
    paragraph.paragraph_format.first_line_indent = Mm(0)


def _bullet(
    document,
    text: object,
    font_size: float,
    *,
    exact_line_height: float | None = None,
    after: float = 1.5,
    fonts: dict | None = None,
    text_indent_mm: float = 4.5,
    native_hanging: bool = False,
    native_marker_gap_mm: float = 0.8,
) -> None:
    value = str(text)
    line_spacing = Pt(exact_line_height) if exact_line_height is not None else 1.12
    marker_match = _NATIVE_LIST_MARKER_RE.match(value)
    if marker_match:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _paragraph_spacing(paragraph, after=after, line=line_spacing)
        if native_hanging:
            _set_numbered_indent(
                paragraph,
                text_indent_mm=text_indent_mm,
                marker_gap_mm=native_marker_gap_mm,
            )
            _set_font(paragraph.add_run("\t"), font_size, fonts=fonts)
            _set_font(paragraph.add_run(marker_match.group(1)), font_size, fonts=fonts)
            _set_font(paragraph.add_run("\t"), font_size, fonts=fonts)
            _add_markdown_runs(paragraph, marker_match.group(2), font_size, fonts=fonts)
        else:
            _set_flush_indent(paragraph)
            _add_markdown_runs(paragraph, value, font_size, fonts=fonts)
        return
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _paragraph_spacing(paragraph, after=after, line=line_spacing)
    _set_hanging_indent(paragraph, left=text_indent_mm, hanging=min(3.0, text_indent_mm))
    _add_markdown_runs(paragraph, value.lstrip("• "), font_size, fonts=fonts)


def generate_docx(
    resume_data: dict,
    style: dict | None = None,
    photo: str | None = None,
    lang: str = "zh",
    layout_config: dict | None = None,
) -> bytes:
    """Return an editable DOCX that follows the controlled layout configuration."""
    from .layout import apply_page_mode_defaults
    from .layout_config import (
        format_compact_academic_metric,
        normalize_layout_config,
        resolve_content_block_flow,
        resolve_education_column_widths,
        resolve_layout_tokens,
        resolve_photo_height_mm,
    )

    data = normalize_resume_data(resume_data)
    labels = LABELS.get(lang, LABELS["zh"])
    colon = "：" if lang == "zh" else ": "
    layout = normalize_layout_config(layout_config)
    global_layout = layout["global"]
    def merged_into_education(section_id: str) -> bool:
        return bool(data.get("education")) and global_layout.get("sectionPlacements", {}).get(section_id) == "education"
    style = apply_page_mode_defaults(style)
    tokens = resolve_layout_tokens(layout, style)
    tokens["photoHeightMm"] = resolve_photo_height_mm(
        data,
        layout,
        tokens,
        photo_present=bool(photo or (data.get("basics") or {}).get("photo")),
    )
    font_spec = {
        "latinFont": tokens["latinFont"],
        "eastAsiaFont": tokens["eastAsiaFont"],
    }
    # Word stores font sizes in half-point units. The shared layout resolver
    # already defines a compact semantic scale, so each export role uses the
    # same physical size as the browser and PDF renderers.
    body_font_size = round(tokens["bodyFontSizePt"] * 2) / 2
    meta_font_size = round(tokens["metaFontSizePt"] * 2) / 2
    entry_title_font_size = round(tokens["entryTitleFontSizePt"] * 2) / 2
    section_title_font_size = round(tokens["sectionTitleFontSizePt"] * 2) / 2
    name_font_size = round(tokens["nameFontSizePt"] * 2) / 2
    label_font_size = round(tokens["labelFontSizePt"] * 2) / 2
    body_bold = tokens["bodyFontWeight"] >= 600
    meta_bold = tokens["metaFontWeight"] >= 600
    entry_title_bold = tokens["entryTitleFontWeight"] >= 600
    section_title_bold = tokens["sectionTitleFontWeight"] >= 600
    name_bold = tokens["nameFontWeight"] >= 600
    label_bold = tokens["labelFontWeight"] >= 600
    formatting_version = int(data.get("formatting_version") or 0)
    legacy_default_bold = formatting_version < 1
    legacy_section_title_bold = formatting_version < 2
    legacy_manual_field_bold = formatting_version < 3

    body_line_height = body_font_size * tokens["lineHeight"]
    printable_width_mm = 210.0 - tokens["marginLeftMm"] - tokens["marginRightMm"]
    list_text_indent_mm = tokens["listTextIndentPt"] * 25.4 / 72.0
    list_marker_gap_mm = tokens["listMarkerGapPt"] * 25.4 / 72.0
    module_spacing = tokens["moduleSpacingPt"]
    page_break_before = style.get("pageBreakBefore", "")
    hidden_sections = set(global_layout["hiddenSections"])

    document = Document()
    _configure_chinese_document(document)
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = Mm(tokens["marginTopMm"])
    section.bottom_margin = Mm(tokens["marginBottomMm"])
    section.left_margin = Mm(tokens["marginLeftMm"])
    section.right_margin = Mm(tokens["marginRightMm"])
    _disable_document_grid(section)

    normal = document.styles["Normal"]
    normal.font.name = font_spec["latinFont"]
    normal_fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("ascii", "hAnsi", "cs"):
        normal_fonts.set(qn(f"w:{script}"), font_spec["latinFont"])
    normal_fonts.set(qn("w:eastAsia"), font_spec["eastAsiaFont"])
    _set_language(normal._element.get_or_add_rPr())
    normal.font.size = Pt(body_font_size)
    normal.font.bold = body_bold
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(body_line_height)
    normal.paragraph_format.widow_control = False

    photo_anchor_added = False
    photo_height_mm = tokens["photoHeightMm"]
    basics_height_pt = 0.0

    def title(section_id: str, fallback: str) -> None:
        nonlocal photo_anchor_added, photo_height_mm, photo_width_mm
        text = global_layout.get("titleOverrides", {}).get(section_id, {}).get(lang)
        if text is None:
            text = f"**{fallback}**" if formatting_version >= 2 else fallback
        module_layout = layout.get(section_id, layout["custom_sections"])
        paragraph = document.add_paragraph()
        paragraph.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[module_layout.get("titleAlignment") or "left"]
        _paragraph_spacing(
            paragraph,
            before=module_spacing,
            after=tokens["sectionTitleAfterPt"],
            line=Pt(tokens["sectionTitleLineHeightPt"]),
        )
        _add_markdown_runs(
            paragraph,
            text,
            section_title_font_size,
            fonts=font_spec,
            base_bold=section_title_bold if legacy_section_title_bold else False,
        )
        if (module_layout.get("titleStyle") or global_layout["titleStyle"]) in {"plain", "underline"}:
            p_pr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (
                ("val", "single"),
                ("sz", "6"),
                ("space", str(round(tokens["sectionTitleBorderGapPt"]))),
                ("color", "333333"),
            ):
                bottom.set(qn(f"w:{key}"), value)
            borders.append(bottom)
            p_pr.append(borders)
        if photo_bytes and not photo_anchor_added:
            # The page/PDF height is the shared visual baseline. Word's own
            # exact table/paragraph geometry can place its divider slightly
            # higher, so cap only the DOCX anchor before that border. This
            # does not participate in document flow and never moves the title.
            photo_height_mm = _resolve_word_photo_height_mm(
                photo_height_mm,
                basics_height_pt=basics_height_pt,
                title_paragraph=paragraph,
            )
            photo_width_mm = photo_height_mm * _photo_aspect_ratio(photo_bytes, data)
            _add_floating_picture(
                paragraph,
                photo_bytes,
                width_mm=photo_width_mm,
                height_mm=photo_height_mm,
                x_mm=tokens["marginLeftMm"] + printable_width_mm - photo_width_mm,
                y_mm=tokens["marginTopMm"],
            )
            photo_anchor_added = True

    def maybe_break(key: str) -> None:
        if page_break_before == key:
            document.add_page_break()

    def add_details(details, details_style: str) -> None:
        for detail in details or []:
            if details_style == "paragraph":
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _paragraph_spacing(paragraph, after=tokens["paragraphSpacingPt"], line=Pt(body_line_height))
                _add_markdown_runs(paragraph, detail, body_font_size, fonts=font_spec)
            else:
                _bullet(
                    document,
                    detail,
                    body_font_size,
                    exact_line_height=body_line_height,
                    after=tokens["paragraphSpacingPt"],
                    fonts=font_spec,
                    text_indent_mm=list_text_indent_mm,
                )

    basics = data.get("basics") or {}
    basics_layout = layout["basics"]
    hidden_basics = set(basics_layout["hiddenFields"])
    contact_values = [str(basics.get(key)) for key in ("gender",) if basics.get(key) and key not in hidden_basics]
    if basics.get("birth_date") and "birth_date" not in hidden_basics:
        contact_values.append(str(basics["birth_date"]))
    direct_contact = [str(basics.get(key)) for key in ("phone", "email") if basics.get(key) and key not in hidden_basics]
    additional_values = [
        (
            f'{item.get("label")}{colon}{item.get("value")}'
            if item.get("label") and item.get("value")
            else str(item.get("label") or item.get("value") or "")
        )
        for item in basics.get("additional_fields", [])
        if (item.get("label") or item.get("value")) and "additional_fields" not in hidden_basics
    ]
    photo_bytes = None
    if (photo or basics.get("photo")) and "photo" not in hidden_basics:
        try:
            encoded = (photo or basics.get("photo")).split(",", 1)[-1]
            photo_bytes = base64.b64decode(encoded)
        except (ValueError, TypeError):
            photo_bytes = None
    photo_width_mm = tokens["photoHeightMm"] * _photo_aspect_ratio(photo_bytes, data)
    basics_values = {
        "name": basics.get("name") or labels["nameNotSet"],
        "target_position": (
            f'{"**" if _is_fully_bold(basics["target_position"]) else ""}'
            f'{labels["targetPosition"]}{colon}'
            f'{"**" if _is_fully_bold(basics["target_position"]) else ""}'
            f'{basics["target_position"]}'
            if basics.get("target_position") and "target_position" not in hidden_basics else ""
        ),
        "personal_meta": " | ".join(contact_values),
        "contact": " | ".join(direct_contact),
        "additional_fields": " | ".join(additional_values),
        "photo": photo_bytes,
    }
    hidden_basic_components = set(basics_layout["hiddenComponents"])
    alignment_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    # Word's basic-information table contains text only.  The photo is a
    # page-positioned anchor added to the first section heading below, so it
    # cannot distort the text table's row/column geometry.
    top_components = []
    meta_components = []
    component_alignment = {}
    for component_row in basics_layout["componentRows"]:
        for cell_config in component_row["cells"]:
            for component in cell_config["components"]:
                if (
                    component == "photo"
                    or component in hidden_basic_components
                    or not basics_values.get(component)
                ):
                    continue
                component_alignment.setdefault(component, alignment_map[cell_config["alignment"]])
                target = meta_components if component in {"personal_meta", "contact", "additional_fields"} else top_components
                if component not in target:
                    target.append(component)

    if top_components or meta_components:
        table_width_mm = printable_width_mm
        if photo_bytes:
            # Keep the text block on the left and leave a deliberate visual
            # gap before the top-right photo, matching the browser/PDF header.
            table_width_mm = max(40.0, printable_width_mm - photo_width_mm - 40.0)
        table = document.add_table(rows=0, cols=1)
        _set_fixed_table_widths(table, (table_width_mm,))
        # Keep each top-level basic-information component in its own row.
        # This produces the same three-row, one-column structure as the
        # browser/PDF header: name, target position, then contact metadata.
        for component in top_components:
            cell = table.add_row().cells[0]
            _set_cell_borderless(cell)
            _set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = component_alignment.get(component, WD_ALIGN_PARAGRAPH.LEFT)
            _paragraph_spacing(
                paragraph,
                after=tokens["modules"]["basics"]["rowSpacingPt"],
                line=Pt(tokens["nameLineHeightPt"] if component == "name" else tokens["metaLineHeightPt"]),
            )
            if component == "name":
                _add_markdown_runs(
                    paragraph,
                    basics_values[component],
                    name_font_size,
                    fonts=font_spec,
                    base_bold=name_bold if legacy_default_bold else False,
                )
            else:
                _add_markdown_runs(
                    paragraph,
                    basics_values[component],
                    meta_font_size,
                    color="333333",
                    fonts=font_spec,
                    base_bold=component == "target_position" and legacy_default_bold,
                )
            line_height_pt = tokens["nameLineHeightPt"] if component == "name" else tokens["metaLineHeightPt"]
            basics_height_pt += (
                _word_estimated_line_count(
                    basics_values[component],
                    name_font_size if component == "name" else meta_font_size,
                    table_width_mm,
                )
                * _word_paragraph_line_height_pt(paragraph, line_height_pt)
                + _word_paragraph_spacing_pt(paragraph, "after")
            )
        if meta_components:
            cell = table.add_row().cells[0]
            _set_cell_borderless(cell)
            _set_cell_margins(cell)
            parts = [basics_values[component] for component in meta_components if basics_values.get(component)]
            paragraph = cell.paragraphs[0]
            paragraph.alignment = component_alignment.get(meta_components[0], WD_ALIGN_PARAGRAPH.LEFT)
            _paragraph_spacing(
                paragraph,
                after=tokens["modules"]["basics"]["rowSpacingPt"],
                line=Pt(tokens["metaLineHeightPt"]),
            )
            _add_markdown_runs(
                paragraph,
                " | ".join(parts),
                meta_font_size,
                color="333333",
                fonts=font_spec,
            )
            basics_height_pt += (
                _word_estimated_line_count(
                    " | ".join(parts),
                    meta_font_size,
                    table_width_mm,
                )
                * _word_paragraph_line_height_pt(paragraph, tokens["metaLineHeightPt"])
                + _word_paragraph_spacing_pt(paragraph, "after")
            )

    def render_education() -> None:
        items = data.get("education") or []
        if not items or "education" in hidden_sections:
            return
        cfg = layout["education"]
        hidden_metrics = set(cfg["hiddenMetrics"])
        schools = [str(item.get("school_name") or labels["schoolNotSet"]) for item in items]
        dates = [_date_range(item) for item in items]
        degree_majors = [
            " · ".join(value for value in (item.get("degree", ""), item.get("major", "")) if value)
            for item in items
        ]

        def compact_metric(item: dict) -> str:
            return format_compact_academic_metric(
                item,
                hidden_metrics,
            )

        compact_metrics = [compact_metric(item) for item in items]
        compact_widths = resolve_education_column_widths(
            tokens,
            schools=schools,
            dates=dates,
            degree_majors=degree_majors,
            compact_metrics=compact_metrics,
        )
        maybe_break("education:0")
        title("education", labels["education"])
        for index, item in enumerate(items):
            item_paragraph_start = len(document.paragraphs)
            item_table_start = len(document.tables)
            if index:
                maybe_break(f"education:{index}")
            school = item.get("school_name") or labels["schoolNotSet"]
            tags = item.get("school_tags") or []
            tag_text = " · ".join(str(tag) for tag in tags)
            if cfg["schoolTagStyle"] == "hidden":
                tag_text = ""
            degree = " · ".join(value for value in (item.get("degree", ""), item.get("major", "")) if value)
            degree = degree_majors[index]
            date = dates[index]
            metrics = []
            if item.get("gpa") and "gpa" not in hidden_metrics:
                value = str(item["gpa"]) + (f'/{item["gpa_scale"]}' if item.get("gpa_scale") else "")
                metrics.append(f'{labels["gpa"]}{colon}{value}')
            if item.get("ranking") and "ranking" not in hidden_metrics:
                metrics.append(f'{labels["ranking"]}{colon}{item["ranking"]}')
            component_values = {
                "school": school,
                "school_tags": tag_text,
                "degree": item.get("degree", ""),
                "major": item.get("major", ""),
                "metrics": compact_metrics[index] if cfg["preset"] == "compact" else " · ".join(metrics),
                "date": date,
            }
            hidden_components = set(cfg["hiddenComponents"]) | {"theses"}
            alignment_map = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            for component_row in cfg["componentRows"]:
                active_cells = []
                for cell_config in component_row["cells"]:
                    components = [
                        component for component in cell_config["components"]
                        if component not in hidden_components and component_values.get(component)
                    ]
                    if components:
                        active_cells.append((cell_config, components))
                if not active_cells:
                    continue
                is_compact_header = (
                    cfg["preset"] == "compact"
                    and len(active_cells) == 3
                    and "school" in active_cells[0][1]
                    and any(component in active_cells[1][1] for component in ("degree", "major", "metrics"))
                    and "date" in active_cells[2][1]
                )
                if is_compact_header:
                    widths = (compact_widths["sideMm"], compact_widths["middleMm"], compact_widths["sideMm"])
                else:
                    content_width = min(42.0, printable_width_mm / max(1, len(active_cells)))
                    fixed_width = sum(content_width for cell, _ in active_cells if cell["width"] == "content")
                    flexible_count = max(1, sum(1 for cell, _ in active_cells if cell["width"] != "content"))
                    flexible_width = max(10.0, (printable_width_mm - fixed_width) / flexible_count)
                    widths = tuple(content_width if cell["width"] == "content" else flexible_width for cell, _ in active_cells)
                table = document.add_table(rows=1, cols=len(active_cells))
                _set_fixed_table_widths(table, widths)
                for cell_index, ((cell_config, components), cell) in enumerate(zip(active_cells, table.rows[0].cells)):
                    _set_cell_borderless(cell)
                    _set_cell_margins(cell)
                    for paragraph_index, component in enumerate(components):
                        paragraph = (
                            cell.paragraphs[0]
                            if paragraph_index == 0 or cell_config["flow"] == "inline"
                            else cell.add_paragraph()
                        )
                        paragraph.alignment = alignment_map[cell_config["alignment"]]
                        _paragraph_spacing(
                            paragraph,
                            after=(tokens["modules"]["education"]["rowSpacingPt"] if paragraph_index == len(components) - 1 else 0),
                            line=Pt(tokens["entryTitleLineHeightPt"]),
                        )
                        prefix = " · " if cell_config["flow"] == "inline" and paragraph_index > 0 else ""
                        if component == "school":
                            _add_markdown_runs(paragraph, prefix + component_values[component], entry_title_font_size, fonts=font_spec, base_bold=entry_title_bold if legacy_default_bold else False)
                        elif component in {"school_tags", "degree", "major", "metrics", "date"}:
                            # These are user-entered field values. Their visual
                            # weight comes only from explicit inline-bold marks.
                            _add_markdown_runs(paragraph, prefix + component_values[component], label_font_size, fonts=font_spec, base_bold=legacy_manual_field_bold)
                        else:
                            _add_markdown_runs(paragraph, prefix + component_values[component], meta_font_size, fonts=font_spec, base_bold=meta_bold)
            if cfg["thesisDisplay"] != "hidden" and "theses" not in cfg["hiddenComponents"]:
                for thesis in item.get("theses") or []:
                    if not isinstance(thesis, dict):
                        continue
                    if thesis.get("title"):
                        p = document.add_paragraph()
                        if label_font_size == body_font_size:
                            _add_markdown_runs(
                                p,
                                f'{labels["thesis"]}{colon}{thesis["title"]}',
                                body_font_size,
                                fonts=font_spec,
                                base_bold=label_bold,
                            )
                        else:
                            _set_font(p.add_run(f'{labels["thesis"]}{colon}'), label_font_size, bold=label_bold, fonts=font_spec)
                            _add_markdown_runs(p, thesis["title"], body_font_size, fonts=font_spec, base_bold=label_bold)
                    if cfg["thesisDisplay"] == "expanded":
                        add_details(thesis.get("details"), "bullets")
            if len(document.paragraphs) > item_paragraph_start:
                _increase_paragraph_after(document.paragraphs[-1], tokens["modules"]["education"]["itemSpacingPt"])
            elif len(document.tables) > item_table_start:
                for cell in document.tables[-1].rows[-1].cells:
                    _increase_paragraph_after(cell.paragraphs[-1], tokens["modules"]["education"]["itemSpacingPt"])

        merged_sections = [
            ("research_interests", labels["researchInterests"], data.get("research_interests") or []),
            ("honors", labels["honors"], data.get("honors") or []),
            ("publications", "Publications" if lang == "en" else "论文", data.get("publications") or []),
        ]
        other_values = data.get("others") or {}
        other_cfg = layout["others"]
        other_hidden = set(other_cfg["hiddenFields"]) | set(other_cfg["hiddenComponents"])
        other_labels = _other_field_labels(other_values, labels)
        separator = " · " if other_cfg["separator"] == "dot" else " | "
        merged_other = [
            _other_field_value(other_labels[field], other_values[field], separator, colon)
            for field in other_cfg["fieldOrder"]
            if field in other_labels and field not in other_hidden and other_values.get(field)
        ]
        merged_sections.append(("others", "Certificates & Languages" if lang == "en" else "证书与语言", merged_other))
        merged_sections.sort(key=lambda item: global_layout["sectionOrder"].index(item[0]))
        for section_id, fallback, values in merged_sections:
            if not merged_into_education(section_id) or section_id in hidden_sections or not values:
                continue
            heading = document.add_paragraph()
            _paragraph_spacing(heading, after=tokens["paragraphSpacingPt"], line=Pt(body_line_height))
            heading_text = global_layout.get("titleOverrides", {}).get(section_id, {}).get(lang)
            if heading_text is None:
                heading_text = f"**{fallback}**" if formatting_version >= 2 else fallback
            _add_markdown_runs(heading, heading_text, body_font_size, fonts=font_spec, base_bold=body_bold)
            for value_index, value in enumerate(values):
                module_id = section_id if section_id in tokens["modules"] and "listStyle" in layout.get(section_id, {}) else "publications"
                add_module_list_item(module_id, value, value_index)

    def add_component_rows(module_id: str, values: dict[str, str], excluded: set[str]) -> None:
        cfg = layout[module_id]
        hidden = set(cfg["hiddenComponents"]) | set(excluded)
        alignment_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        title_components = {"organization", "project_name", "school"}
        label_components = {"school_tags", "degree", "major", "metrics", "date", "position", "job_type"}
        for component_row in cfg["componentRows"]:
            active_cells = []
            for cell_cfg in component_row["cells"]:
                components = [item for item in cell_cfg["components"] if item not in hidden and values.get(item)]
                if components:
                    active_cells.append((cell_cfg, components))
            if not active_cells:
                continue
            content_width = min(42.0, printable_width_mm / max(1, len(active_cells)))
            fixed = sum(content_width for cell_cfg, _ in active_cells if cell_cfg["width"] == "content")
            flexible_count = max(1, sum(1 for cell_cfg, _ in active_cells if cell_cfg["width"] != "content"))
            flexible = max(10.0, (printable_width_mm - fixed) / flexible_count)
            widths = tuple(content_width if cell_cfg["width"] == "content" else flexible for cell_cfg, _ in active_cells)
            table = document.add_table(rows=1, cols=len(active_cells))
            _set_fixed_table_widths(table, widths)
            for (cell_cfg, components), cell in zip(active_cells, table.rows[0].cells):
                _set_cell_borderless(cell)
                _set_cell_margins(cell)
                for component_index, component in enumerate(components):
                    paragraph = cell.paragraphs[0] if component_index == 0 or cell_cfg["flow"] == "inline" else cell.add_paragraph()
                    paragraph.alignment = alignment_map[cell_cfg["alignment"]]
                    _paragraph_spacing(paragraph, after=tokens["modules"][module_id]["rowSpacingPt"], line=Pt(tokens["entryTitleLineHeightPt"]))
                    prefix = (
                        " " if cell_cfg["flow"] == "inline" and component_index and component == "job_type"
                        else (" · " if cell_cfg["flow"] == "inline" and component_index else "")
                    )
                    size = entry_title_font_size if component in title_components else (label_font_size if component in label_components else meta_font_size)
                    bold = (entry_title_bold if legacy_default_bold else False) if component in title_components else (legacy_manual_field_bold if component in label_components else meta_bold)
                    _add_markdown_runs(paragraph, prefix + values[component], size, fonts=font_spec, base_bold=bold)

    def add_module_list_item(module_id: str, value: object, index: int) -> None:
        cfg = layout[module_id]
        module_tokens = tokens["modules"][module_id]
        base_indent_mm = module_tokens["indentPt"] * 25.4 / 72.0
        after = module_tokens["itemSpacingPt"]
        list_style = cfg.get("listStyle", "bullet")
        content = _module_list_content(value)
        marker_bold = _is_fully_bold(value)
        if list_style == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.left_indent = Mm(base_indent_mm)
            _paragraph_spacing(paragraph, after=after, line=Pt(body_line_height))
            _add_markdown_runs(paragraph, content, body_font_size, fonts=font_spec)
            return
        text_indent_mm = base_indent_mm + list_text_indent_mm
        if list_style == "bullet":
            _bullet(
                document,
                content,
                body_font_size,
                exact_line_height=body_line_height,
                after=after,
                fonts=font_spec,
                text_indent_mm=text_indent_mm,
                native_hanging=True,
                native_marker_gap_mm=list_marker_gap_mm,
            )
            return
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _paragraph_spacing(paragraph, after=after, line=Pt(body_line_height))
        _set_numbered_indent(
            paragraph,
            text_indent_mm=text_indent_mm,
            marker_gap_mm=list_marker_gap_mm,
            extra_indent_mm=WORD_NUMBERED_EXTRA_INDENT_MM,
        )
        _set_font(paragraph.add_run("\t"), body_font_size, fonts=font_spec)
        _set_font(paragraph.add_run(f"({index + 1})"), body_font_size, bold=marker_bold, fonts=font_spec)
        _set_font(paragraph.add_run("\t"), body_font_size, fonts=font_spec)
        _add_markdown_runs(paragraph, content, body_font_size, fonts=font_spec)

    def work_groups():
        items = list(data.get("work_experience") or [])
        if not global_layout["splitWorkExperience"]:
            return [("work_experience", labels["workExperience"], items)]
        regular = [item for item in items if not re.search(r"实习|intern", str(item.get("job_type", "")), re.I)]
        interns = [item for item in items if re.search(r"实习|intern", str(item.get("job_type", "")), re.I)]
        return [("work_experience", labels["workExperience"], regular), ("internship_experience", "Internship Experience" if lang == "en" else "实习经历", interns)]

    def render_work(section_id: str, fallback: str, items: list) -> None:
        if not items or section_id in hidden_sections:
            return
        cfg = layout[section_id]
        maybe_break(f"{section_id}:0")
        title(section_id, fallback)
        for index, item in enumerate(items):
            if index:
                maybe_break(f"{section_id}:{index}")
            company = item.get("company_name", "")
            add_component_rows(section_id, {
                "organization": company,
                "position": item.get("job_title", ""),
                "job_type": f'({item["job_type"]})' if cfg["showJobType"] and item.get("job_type") else "",
                "date": _date_range(item),
            }, {"content"})
            detail_start = len(document.paragraphs)
            add_content_blocks(item, section_id)
            if len(document.paragraphs) > detail_start:
                _increase_paragraph_after(document.paragraphs[-1], tokens["modules"][section_id]["itemSpacingPt"])
            else:
                for cell in document.tables[-1].rows[-1].cells:
                    _increase_paragraph_after(cell.paragraphs[-1], tokens["modules"][section_id]["itemSpacingPt"])

    def add_content_blocks(item: dict, module_id: str) -> None:
        module_tokens = tokens["modules"][module_id]
        module_indent_mm = module_tokens["indentPt"] * 25.4 / 72.0
        semantic_indent_mm = module_indent_mm + list_text_indent_mm
        blocks = item.get("content_blocks") or []
        if not blocks:
            add_details(item.get("details"), "bullets")
            return
        for block in blocks:
            flow = resolve_content_block_flow(block)
            if not flow["visible"]:
                continue
            block_type = flow["type"]
            label = flow["label"]
            if block_type == "paragraph":
                paragraph = document.add_paragraph(style="List Bullet") if flow["labelMarker"] == "bullet" else document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _paragraph_spacing(paragraph, after=module_tokens["contentBlockSpacingPt"], line=Pt(body_line_height))
                if flow["labelMarker"] == "bullet":
                    _set_hanging_indent(paragraph, left=semantic_indent_mm, hanging=min(3.0, list_text_indent_mm))
                else:
                    paragraph.paragraph_format.left_indent = Mm(module_indent_mm)
                if label:
                    _add_markdown_runs(paragraph, f"{label}{colon}", body_font_size, fonts=font_spec, base_bold=label_bold and flow["labelBold"])
                _add_markdown_runs(paragraph, block.get("text", ""), body_font_size, fonts=font_spec)
                continue
            if flow["labelPlacement"] == "separate":
                paragraph = document.add_paragraph(style="List Bullet") if flow["labelMarker"] == "bullet" else document.add_paragraph()
                _paragraph_spacing(paragraph, after=tokens["contentLabelSpacingPt"], line=Pt(body_line_height))
                if flow["labelMarker"] == "bullet":
                    _set_hanging_indent(paragraph, left=semantic_indent_mm, hanging=min(3.0, list_text_indent_mm))
                _add_markdown_runs(paragraph, f"{label}{colon}", body_font_size, fonts=font_spec, base_bold=label_bold and flow["labelBold"])
            details = block.get("items") or []
            for detail_index, detail in enumerate(details, 1):
                trailing_block_spacing = module_tokens["contentBlockSpacingPt"] if detail_index == len(details) else 0
                if block_type == "numbered_list":
                    paragraph = document.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    _paragraph_spacing(
                        paragraph,
                        after=tokens["numberedItemSpacingPt"] + trailing_block_spacing,
                        line=Pt(body_line_height),
                    )
                    _set_numbered_indent(
                        paragraph,
                        text_indent_mm=module_indent_mm + list_text_indent_mm * max(1, flow["contentIndentLevels"]),
                        marker_gap_mm=list_marker_gap_mm,
                    )
                    _set_font(paragraph.add_run("\t"), body_font_size, fonts=font_spec)
                    _set_font(paragraph.add_run(f"({detail_index})"), body_font_size, bold=_is_fully_bold(detail), fonts=font_spec)
                    _set_font(paragraph.add_run("\t"), body_font_size, fonts=font_spec)
                    _add_markdown_runs(paragraph, detail, body_font_size, fonts=font_spec)
                else:
                    _bullet(
                        document,
                        detail,
                        body_font_size,
                        exact_line_height=body_line_height,
                        after=module_tokens["paragraphSpacingPt"] + trailing_block_spacing,
                        fonts=font_spec,
                        text_indent_mm=module_indent_mm + list_text_indent_mm * max(1, flow["contentIndentLevels"]),
                    )

    def render_projects() -> None:
        items = data.get("project_experience") or []
        if not items or "project_experience" in hidden_sections:
            return
        cfg = layout["project_experience"]
        maybe_break("project_experience:0")
        title("project_experience", labels["projectExperience"])
        for index, item in enumerate(items):
            if index:
                maybe_break(f"project_experience:{index}")
            project_name = item.get("project_name") or item.get("name") or labels["projectNotSet"]
            role = item.get("role") if cfg["showRole"] else ""
            date = _date_range(item) if cfg["showDate"] else ""
            add_component_rows("project_experience", {
                "project_name": project_name,
                "role": role,
                "date": date,
            }, {"content"})
            detail_start = len(document.paragraphs)
            add_content_blocks(item, "project_experience")
            if len(document.paragraphs) > detail_start:
                _increase_paragraph_after(document.paragraphs[-1], tokens["modules"]["project_experience"]["itemSpacingPt"])
            else:
                for cell in document.tables[-1].rows[-1].cells:
                    _increase_paragraph_after(cell.paragraphs[-1], tokens["modules"]["project_experience"]["itemSpacingPt"])

    def render_skills() -> None:
        skills = (data.get("others") or {}).get("skills") or []
        if not skills or "skills" in hidden_sections:
            return
        title("skills", labels["skills"])
        for index, value in enumerate(skills):
            add_module_list_item("skills", value, index)

    def render_others() -> None:
        values = data.get("others") or {}
        cfg = layout["others"]
        hidden_fields = set(cfg["hiddenFields"]) | set(cfg["hiddenComponents"])
        fields = [key for key in cfg["fieldOrder"] if key != "skills" and key not in hidden_fields and values.get(key)]
        if not fields or "others" in hidden_sections or merged_into_education("others"):
            return
        maybe_break("others")
        title("others", "Certificates & Languages" if lang == "en" else "证书与语言")
        field_labels = {"skills": labels["skills"], **_other_field_labels(values, labels)}
        separator = " · " if cfg["separator"] == "dot" else " | "
        component_values = {
            key: _other_field_value(field_labels[key], values[key], separator, colon)
            for key in fields
        }
        add_component_rows("others", component_values, set())

    def render_self() -> None:
        values = data.get("self_evaluation") or []
        if not values or "self_evaluation" in hidden_sections:
            return
        cfg = layout["self_evaluation"]
        maybe_break("self_evaluation")
        title("self_evaluation", labels["selfEvaluation"])
        if cfg["preset"] == "compact":
            values = [" ".join(str(value) for value in values)]
        for index, value in enumerate(values):
            add_module_list_item("self_evaluation", value, index)

    def render_plain_section(section_id: str, fallback: str, values: list) -> None:
        if not values or section_id in hidden_sections or merged_into_education(section_id):
            return
        maybe_break(section_id)
        title(section_id, fallback)
        for index, value in enumerate(values):
            add_module_list_item(section_id, value, index)

    def render_research() -> None:
        render_plain_section("research_interests", labels["researchInterests"], data.get("research_interests") or [])

    def render_honors() -> None:
        render_plain_section("honors", labels["honors"], data.get("honors") or [])

    def render_publications() -> None:
        render_plain_section("publications", "Publications" if lang == "en" else "论文", data.get("publications") or [])

    def render_custom_sections() -> None:
        if "custom_sections" in hidden_sections:
            return
        for index, custom in enumerate(data.get("custom_sections") or []):
            if not custom.get("title") or not custom.get("items"):
                continue
            maybe_break(f"custom_sections:{index}")
            title("custom_sections", custom["title"])
            for item_index, value in enumerate(custom["items"]):
                add_module_list_item("custom_sections", value, item_index)

    renderers = {
        "education": render_education,
        "skills": render_skills,
        "research_interests": render_research,
        "honors": render_honors,
        "publications": render_publications,
        "project_experience": render_projects,
        "custom_sections": render_custom_sections,
        "others": render_others,
        "self_evaluation": render_self,
    }
    work_by_id = {section_id: (fallback, items) for section_id, fallback, items in work_groups()}
    for section_id in global_layout["sectionOrder"]:
        if section_id in work_by_id:
            fallback, items = work_by_id[section_id]
            render_work(section_id, fallback, items)
        elif section_id in renderers:
            renderers[section_id]()

    if photo_bytes and not photo_anchor_added:
        paragraph = document.add_paragraph()
        _paragraph_spacing(paragraph, after=0, line=Pt(1))
        _add_floating_picture(
            paragraph,
            photo_bytes,
            width_mm=photo_width_mm,
            height_mm=photo_height_mm,
            x_mm=tokens["marginLeftMm"] + printable_width_mm - photo_width_mm,
            y_mm=tokens["marginTopMm"],
        )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
