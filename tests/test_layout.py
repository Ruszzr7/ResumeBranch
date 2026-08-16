import unittest
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from backend.layout import (
    apply_page_mode_defaults,
    enforce_page_limit,
    normalize_page_break,
    normalize_page_mode,
    normalize_source_page_count,
    page_limit,
)
from backend.docx_generator import generate_docx
from backend.pdf_generator import render_resume_to_html
from backend.layout_config import default_layout_config, normalize_layout_config, resolve_layout_tokens


def resume_with_two_jobs():
    return {
        "basics": {"name": "测试", "gender": "", "phone": "", "email": "", "target_position": ""},
        "education": [],
        "work_experience": [
            {"company_name": "甲", "job_title": "开发", "date_range": ["2024", "2025"], "job_type": "", "details": ["内容"]},
            {"company_name": "乙", "job_title": "开发", "date_range": ["2025", "至今"], "job_type": "", "details": ["内容"]},
        ],
        "project_experience": [],
        "others": {"skills": [], "certificates": [], "languages": []},
        "self_evaluation": [],
    }


class LayoutRuleTests(unittest.TestCase):
    def test_education_merge_falls_back_to_standalone_when_education_is_empty(self):
        data = resume_with_two_jobs()
        data["publications"] = ["Paper title"]
        layout = default_layout_config()
        layout["global"]["sectionPlacements"] = {"publications": "education"}

        html = render_resume_to_html(data, layout_config=layout)
        self.assertIn('class="section-title title-underline" style="text-align:left">', html)
        self.assertIn("Paper title", html)

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Paper title", text)

    def test_column_settings_and_education_merges_match_html_and_word(self):
        data = resume_with_two_jobs()
        data.update({
            "formatting_version": 1,
            "education": [{
                "school_name": "示例大学", "school_tags": ["211", "双一流"],
                "degree": "硕士", "major": "电子信息", "date_range": ["2024", "2027"],
                "theses": [],
            }],
            "research_interests": ["人机协作"],
            "honors": ["一等奖学金"],
            "publications": ["论文标题（中科院一区 Top，IF 10），已接收"],
            "others": {"skills": ["Python", "Linux"], "certificates": ["软件设计师"], "languages": ["英语 CET-6"]},
        })
        layout = default_layout_config()
        layout["global"]["titleOverrides"] = {"skills": {"zh": "技术栈"}}
        layout["global"]["sectionPlacements"] = {
            "research_interests": "education", "honors": "education",
            "publications": "education", "others": "education",
        }
        layout["global"]["sectionOrder"] = [
            "education", "publications", "research_interests", "honors", "skills",
            "work_experience", "project_experience", "custom_sections", "others", "self_evaluation",
        ]
        layout["skills"]["listStyle"] = "numbered"

        html = render_resume_to_html(data, layout_config=layout)
        self.assertIn("技术栈", html)
        self.assertIn('list-style-numbered', html)
        self.assertIn('component-school_tags">211 · 双一流</span>', html)
        self.assertNotIn('<h4 class="subfield-title education-merged-title"', html)
        self.assertNotIn('class="section-title title-underline" style="text-align:left">论文</h2>', html)
        self.assertIn("论文标题（中科院一区 Top，IF 10），已接收", html)
        self.assertIn('<ul class="list-items module-list list-style-bullet">', html)

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("示例大学 · 211 · 双一流", table_text)
        self.assertIn("技术栈", text)
        self.assertIn("论文标题（中科院一区 Top，IF 10），已接收", text)
        paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertNotIn("论文", paragraph_texts)
        self.assertNotIn("研究方向", paragraph_texts)
        self.assertNotIn("主要荣誉", paragraph_texts)
        self.assertNotIn("证书与语言", paragraph_texts)

    def test_certificates_and_languages_render_as_a_standalone_ordered_section(self):
        data = resume_with_two_jobs()
        data["others"] = {
            "skills": [],
            "certificates": ["软件设计师"],
            "languages": ["英语 CET-6"],
        }
        layout = default_layout_config()

        html = render_resume_to_html(data, layout_config=layout)
        self.assertIn("证书与语言", html)
        self.assertIn("证书：软件设计师", html)
        self.assertIn("语言：英语 CET-6", html)

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("证书与语言", text)
        self.assertIn("证书：软件设计师", table_text)
        self.assertIn("语言：英语 CET-6", table_text)

    def test_custom_certificate_and_language_labels_are_shared_and_can_be_blank(self):
        data = resume_with_two_jobs()
        data["others"] = {
            "skills": [],
            "certificates": ["软件设计师"],
            "languages": ["英语 CET-6"],
            "field_labels": {"certificates": "资格证书", "languages": ""},
        }

        html = render_resume_to_html(data, layout_config=default_layout_config())
        self.assertIn("资格证书：软件设计师", html)
        self.assertIn("英语 CET-6", html)
        self.assertNotIn("语言：英语 CET-6", html)
        self.assertLess(html.index("证书与语言"), html.index("资格证书：软件设计师"))

        document = Document(BytesIO(generate_docx(data, layout_config=default_layout_config())))
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("资格证书：软件设计师", table_text)
        self.assertIn("英语 CET-6", table_text)
        self.assertNotIn("语言：英语 CET-6", table_text)

    def test_explicit_title_formatting_can_cancel_legacy_default_bold_in_all_exports(self):
        data = resume_with_two_jobs()
        data["education"] = [{"school_name": "示例大学", "date_range": [], "school_tags": [], "theses": []}]
        legacy_html = render_resume_to_html(data, layout_config=default_layout_config())
        self.assertIn("--manual-title-font-weight: 700", legacy_html)
        legacy_doc = Document(BytesIO(generate_docx(data, layout_config=default_layout_config())))
        legacy_school = next(cell for table in legacy_doc.tables for row in table.rows for cell in row.cells if "示例大学" in cell.text)
        self.assertTrue(bool(legacy_school.paragraphs[0].runs[0].bold))

        data["formatting_version"] = 1
        explicit_html = render_resume_to_html(data, layout_config=default_layout_config())
        self.assertIn("--manual-title-font-weight: 400", explicit_html)
        explicit_doc = Document(BytesIO(generate_docx(data, layout_config=default_layout_config())))
        explicit_school = next(cell for table in explicit_doc.tables for row in table.rows for cell in row.cells if "示例大学" in cell.text)
        self.assertFalse(bool(explicit_school.paragraphs[0].runs[0].bold))

    def test_v2_module_titles_default_bold_but_allow_explicit_unbold(self):
        data = resume_with_two_jobs()
        data["formatting_version"] = 2
        layout = default_layout_config()

        default_html = render_resume_to_html(data, layout_config=layout)
        self.assertIn("--section-title-font-weight: 400", default_html)
        self.assertIn("<strong>工作经历</strong>", default_html)
        default_doc = Document(BytesIO(generate_docx(data, layout_config=layout)))
        default_title = next(paragraph for paragraph in default_doc.paragraphs if paragraph.text == "工作经历")
        self.assertTrue(bool(default_title.runs[0].bold))

        layout["global"]["titleOverrides"] = {"work_experience": {"zh": "工作经历"}}
        plain_html = render_resume_to_html(data, layout_config=layout)
        self.assertNotIn("<strong>工作经历</strong>", plain_html)
        plain_doc = Document(BytesIO(generate_docx(data, layout_config=layout)))
        plain_title = next(paragraph for paragraph in plain_doc.paragraphs if paragraph.text == "工作经历")
        self.assertFalse(bool(plain_title.runs[0].bold))

    def test_v1_default_layout_migrates_to_compact_high_density_defaults(self):
        layout = normalize_layout_config({
            "version": 1,
            "global": {
                "fontSize": 11, "lineHeight": 1.6, "moduleMargin": 1,
                "marginVertical": 9,
                "sectionOrder": ["education", "project_experience", "others"],
            },
        })
        self.assertEqual(layout["version"], 9)
        self.assertEqual(layout["global"]["fontSize"], 9)
        self.assertEqual(layout["global"]["lineHeight"], 1.28)
        self.assertIn("skills", layout["global"]["sectionOrder"])

    def test_legacy_page_modes_are_normalized_to_automatic(self):
        self.assertEqual(normalize_page_mode("three"), "auto")
        self.assertEqual(normalize_page_mode("one"), "auto")
        self.assertEqual(normalize_page_mode("two"), "auto")
        self.assertEqual(page_limit("one"), 2)
        self.assertEqual(page_limit("two"), 2)
        self.assertEqual(page_limit("auto"), 2)

    def test_layout_keeps_user_spacing_and_normalizes_source_metadata(self):
        style = apply_page_mode_defaults({
            "pageMode": "one", "fontSize": 12, "lineHeight": 1.8,
            "moduleMargin": 1.5, "marginTop": 10, "marginBottom": 10,
            "sourcePageCount": "2", "pageBreakBefore": "work_experience:1",
        })
        self.assertEqual(style["pageMode"], "auto")
        self.assertEqual(style["fontSize"], 12)
        self.assertEqual(style["lineHeight"], 1.8)
        self.assertEqual(style["sourcePageCount"], 2)
        self.assertEqual(style["pageBreakBefore"], "work_experience:1")

    def test_invalid_page_metadata_is_safely_ignored(self):
        self.assertEqual(normalize_source_page_count("bad"), 1)
        self.assertEqual(normalize_page_break("../../bad"), "")
        self.assertEqual(normalize_page_break("education:0"), "education:0")

    def test_three_page_export_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "超出两页"):
            enforce_page_limit("auto", 3)
        enforce_page_limit("one", 2)

    def test_pdf_and_word_apply_the_same_semantic_break_anchor(self):
        style = {"pageBreakBefore": "work_experience:1"}
        html = render_resume_to_html(resume_with_two_jobs(), style)
        self.assertIn(
            'class="work-item preset-compact date-right page-break-before"',
            html,
        )

        docx = generate_docx(resume_with_two_jobs(), style)
        with ZipFile(BytesIO(docx)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:type="page"', document_xml)

    def test_pdf_applies_controlled_layout_presets(self):
        data = resume_with_two_jobs()
        data["education"] = [{
            "school_name": "示例大学", "school_tags": ["211"], "degree": "本科",
            "major": "计算机", "date_range": ["2020", "2024"],
            "gpa": "3.8", "gpa_scale": "4.0", "theses": [],
        }]
        layout = default_layout_config()
        layout["basics"]["preset"] = "left-aligned"
        layout["education"]["preset"] = "three-column"
        layout["education"]["schoolTagStyle"] = "text"
        layout["work_experience"]["detailsStyle"] = "paragraph"
        layout["global"]["titleStyle"] = "plain"
        html = render_resume_to_html(data, layout_config=layout)
        self.assertIn("basics-left-aligned", html)
        self.assertIn("preset-three-column", html)
        header_start = html.index('<div class="module-component-rows">', html.index('class="education-item'))
        header_end = html.index('</div></div></div>', header_start)
        header = html[header_start:header_end]
        self.assertIn('component-degree', header)
        self.assertIn('component-metrics', header)
        self.assertIn('GPA：3.8/4.0', header)
        self.assertIn('component-school_tags', header)
        self.assertIn("details-paragraph", html)
        self.assertIn("title-plain", html)
        self.assertIn("border-bottom: 1px solid #333333", html)

        with ZipFile(BytesIO(generate_docx(data, layout_config=layout))) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:color="333333"', document_xml)

    def test_default_compact_education_restores_symmetric_columns_and_unlabeled_metrics(self):
        data = resume_with_two_jobs()
        data["education"] = [{
            "school_name": "示例大学", "degree": "硕士", "major": "电子信息",
            "date_range": ["2022", "2025"], "gpa": "3.8", "gpa_scale": "5.0",
            "ranking": "前10%", "theses": [],
        }]
        html = render_resume_to_html(data, layout_config=default_layout_config())
        self.assertIn(
            "grid-template-columns:var(--education-compact-side-column) "
            "var(--education-middle-column) var(--education-compact-side-column)",
            html,
        )
        self.assertIn('style="text-align:left;justify-content:flex-start"><span class="module-component component-degree">硕士</span>', html)
        self.assertIn('component-metrics">3.8/5.0 (前10%)</span>', html)
        self.assertNotIn('component-metrics">GPA', html)
        self.assertIn('text-align:right;justify-content:flex-end', html)

    def test_default_layout_keeps_dates_right_aligned_and_work_heading_on_one_row(self):
        data = resume_with_two_jobs()
        data["work_experience"][0].update({"company_name": "示例科技", "job_title": "机器人算法工程师", "job_type": "实习"})
        data["project_experience"] = [{
            "project_name": "协作臂轨迹跟踪系统",
            "role": "",
            "date_range": ["2026.01", "至今"],
            "details": [],
        }]

        layout = default_layout_config()
        tokens = resolve_layout_tokens(layout)
        html = render_resume_to_html(data, layout_config=layout)
        self.assertIn('class="work-item preset-', html)
        self.assertIn('date-right', html)
        self.assertIn('.project-item.date-right .project-header', html)
        self.assertIn('grid-template-columns: minmax(0, 1fr) auto', html)
        self.assertIn('page-break-inside: avoid', html)
        work_start = html.index('<div class="module-component-rows">', html.index('class="work-item'))
        work_end = html.index('</div></div></div>', work_start)
        work_fragment = html[work_start:work_end]
        self.assertIn('示例科技', work_fragment)
        self.assertIn('机器人算法工程师', work_fragment)
        self.assertIn('(实习)', work_fragment)
        self.assertIn('--meta-font-weight: 400;', html)
        self.assertIn('--entry-title-font-weight: 700;', html)
        self.assertIn('font-weight: var(--entry-title-font-weight)', html)
        self.assertIn('font-weight: var(--meta-font-weight)', html)

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        heading_table = next(
            table for table in document.tables
            if any("示例科技" in cell.text for row in table.rows for cell in row.cells)
        )
        heading_cell, date_cell = heading_table.rows[0].cells
        self.assertIn("示例科技 · 机器人算法工程师 (实习)", heading_cell.text)
        self.assertEqual(heading_cell.paragraphs[0].runs[0].text, "示例科技")
        self.assertTrue(heading_cell.paragraphs[0].runs[0].bold)
        self.assertAlmostEqual(heading_cell.paragraphs[0].runs[0].font.size.pt, tokens["entryTitleFontSizePt"], places=1)
        self.assertEqual("".join(run.text for run in heading_cell.paragraphs[0].runs[1:]), " · 机器人算法工程师 (实习)")
        self.assertTrue(all(bool(run.bold) for run in heading_cell.paragraphs[0].runs[1:]))
        self.assertTrue(all(abs(run.font.size.pt - tokens["labelFontSizePt"]) < 0.1 for run in heading_cell.paragraphs[0].runs[1:]))
        self.assertTrue(bool(date_cell.paragraphs[0].runs[0].bold))
        self.assertAlmostEqual(date_cell.paragraphs[0].runs[0].font.size.pt, tokens["labelFontSizePt"], places=1)

    def test_empty_self_evaluation_is_hidden_in_default_layout(self):
        data = resume_with_two_jobs()
        data["self_evaluation"] = ["", "   "]
        html = render_resume_to_html(data, layout_config=default_layout_config())
        self.assertNotIn('<section class="section self-evaluation', html)

    def test_word_applies_section_order_and_hidden_sections(self):
        data = resume_with_two_jobs()
        data["project_experience"] = [{
            "project_name": "先展示项目", "role": "开发", "date_range": [], "details": []
        }]
        data["self_evaluation"] = ["不应出现"]
        layout = default_layout_config()
        layout["global"]["sectionOrder"] = [
            "project_experience", "work_experience", "education", "others", "self_evaluation"
        ]
        layout["global"]["hiddenSections"] = ["self_evaluation"]
        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        combined = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        self.assertNotIn("不应出现", combined)
        with ZipFile(BytesIO(generate_docx(data, layout_config=layout))) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        self.assertLess(xml.index("先展示项目"), xml.index("甲"))


    def test_pdf_and_word_render_lossless_import_fields(self):
        data = resume_with_two_jobs()
        data["basics"]["birth_date"] = "2002.06"
        data["research_interests"] = ["协作臂模型预测阻抗控制"]
        data["honors"] = ["研究生二等奖学金"]
        data["custom_sections"] = [{"title": "校园经历", "items": ["学生组织负责人"]}]

        html = render_resume_to_html(data)

        self.assertIn(">2002.06<", html)
        self.assertNotIn("出生年月：", html)
        self.assertIn("协作臂模型预测阻抗控制", html)
        self.assertIn("研究生二等奖学金", html)
        self.assertIn("校园经历", html)
        self.assertIn("display: block", html)
        self.assertLess(html.index("研究生二等奖学金"), html.index("协作臂模型预测阻抗控制"))

        document = Document(BytesIO(generate_docx(data)))
        combined = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        for expected in ("2002.06", "协作臂模型预测阻抗控制", "研究生二等奖学金", "校园经历", "学生组织负责人"):
            self.assertIn(expected, combined)

    def test_basic_information_uses_pipe_separators_in_all_exports(self):
        data = resume_with_two_jobs()
        data["basics"].update({
            "gender": "男",
            "birth_date": "2002.06",
            "phone": "17622312238",
            "email": "2776553477@qq.com",
        })

        html = render_resume_to_html(data)
        self.assertIn("男 | 2002.06", html)
        self.assertIn("17622312238 | 2776553477@qq.com", html)
        self.assertIn("content: ' | '", html)
        self.assertNotIn("出生年月：", html)

        document = Document(BytesIO(generate_docx(data)))
        combined = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("男 | 2002.06 | 17622312238 | 2776553477@qq.com", combined)

    def test_basic_fields_and_target_label_preserve_manual_bold_in_all_exports(self):
        data = resume_with_two_jobs()
        data["formatting_version"] = 2
        data["basics"].update({
            "gender": "**男**",
            "birth_date": "**2002.06**",
            "target_position": "**软件工程师**",
        })

        html = render_resume_to_html(data)
        self.assertIn("<strong>男</strong> | <strong>2002.06</strong>", html)
        self.assertIn("<strong>目标岗位：软件工程师</strong>", html)
        self.assertNotIn("component-target_position { font-weight: var(--manual-title-font-weight); }", html)

        document = Document(BytesIO(generate_docx(data)))
        target_paragraph = next(
            paragraph
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            if "目标岗位：软件工程师" in paragraph.text
        )
        self.assertTrue(all(run.bold for run in target_paragraph.runs if run.text))

        data["basics"]["target_position"] = "软件工程师"
        plain_document = Document(BytesIO(generate_docx(data)))
        plain_target = next(
            paragraph
            for table in plain_document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            if "目标岗位：软件工程师" in paragraph.text
        )
        self.assertTrue(all(not run.bold for run in plain_target.runs if run.text))

    def test_project_semantic_labels_use_outer_bullets_and_numbered_children(self):
        data = resume_with_two_jobs()
        data["project_experience"] = [{
            "project_name": "机器人控制",
            "role": "",
            "date_range": ["2025.01", "2025.06"],
            "details": ["项目简介：面向展厅导航", "项目职责：", "（1）训练策略", "（2）验证性能"],
        }]
        data["others"]["skills"] = ["ROS2", "Python"]

        html = render_resume_to_html(data)

        self.assertIn('<span class="project-inline-label is-bold">项目简介：</span>面向展厅导航', html)
        self.assertIn('block-paragraph has-semantic-label', html)
        self.assertIn('block-numbered_list has-semantic-label', html)
        self.assertIn('<ol class="project-numbered-list">', html)
        self.assertIn('.project-content-block.has-semantic-label > .project-numbered-list', html)
        self.assertIn('padding-left: var(--list-text-indent)', html)
        self.assertNotIn("（1）训练策略", html)
        self.assertNotIn(">角色<", html)
        self.assertIn("专业技能", html)
        self.assertLess(html.index("专业技能"), html.index("项目经历"))
        self.assertIn("body, .degree-major", html)

    def test_semantic_content_blocks_support_paragraph_bullets_and_numbers(self):
        data = resume_with_two_jobs()
        data["project_experience"] = [{
            "project_name": "机器人控制",
            "content_blocks": [
                {"type": "paragraph", "semantic_role": "introduction", "label": "项目简介", "text": "段落内容", "items": []},
                {"type": "bullet_list", "semantic_role": "responsibilities", "label": "项目职责", "text": "", "items": ["分点内容"]},
                {"type": "numbered_list", "semantic_role": "responsibilities", "label": "实施步骤", "text": "", "items": ["编号内容"]},
            ],
        }]

        html = render_resume_to_html(data)
        self.assertIn("block-paragraph", html)
        self.assertIn("block-bullet_list", html)
        self.assertIn("block-numbered_list", html)
        self.assertIn("段落内容", html)
        self.assertIn("分点内容", html)
        self.assertIn("编号内容", html)
        self.assertIn(".project-inline-label", html)
        self.assertIn("font-size: var(--body-font-size)", html)

        document = Document(BytesIO(generate_docx(data)))
        combined = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for expected in ("段落内容", "分点内容", "编号内容"):
            self.assertIn(expected, combined)

    def test_forced_item_break_preserves_template_and_date_classes(self):
        data = resume_with_two_jobs()
        data["project_experience"] = [
            {"project_name": "A", "date_range": ["2025.01", "2025.02"]},
            {"project_name": "B", "date_range": ["2025.03", "2025.04"]},
        ]
        html = render_resume_to_html(
            data,
            {"pageBreakBefore": "project_experience:1"},
            layout_config=default_layout_config(),
        )
        self.assertIn(
            'class="project-item preset-compact date-right page-break-before"',
            html,
        )

    def test_blank_semantic_label_hides_preserved_content_but_generic_content_remains(self):
        data = resume_with_two_jobs()
        data["work_experience"][0]["content_blocks"] = [
            {
                "type": "paragraph", "semantic_role": "introduction", "label": "",
                "label_bold": True, "text": "保留但隐藏的简介", "items": [],
            },
            {
                "type": "numbered_list", "semantic_role": "responsibilities", "label": "",
                "label_bold": True, "text": "", "items": ["保留但隐藏的职责"],
            },
            {
                "type": "bullet_list", "semantic_role": "generic", "label": "",
                "label_bold": True, "text": "", "items": ["仍然显示的普通内容"],
            },
        ]

        html = render_resume_to_html(data)
        self.assertNotIn("保留但隐藏的简介", html)
        self.assertNotIn("保留但隐藏的职责", html)
        self.assertIn("仍然显示的普通内容", html)

        document = Document(BytesIO(generate_docx(data)))
        combined = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("保留但隐藏的简介", combined)
        self.assertNotIn("保留但隐藏的职责", combined)
        self.assertIn("仍然显示的普通内容", combined)

    def test_semantic_label_weight_and_body_justification_match_pdf_and_word(self):
        data = resume_with_two_jobs()
        data["project_experience"] = [{
            "project_name": "机器人控制",
            "date_range": ["2025.01", "2025.06"],
            "content_blocks": [{
                "type": "paragraph", "label": "项目简介", "label_bold": False,
                "text": "面向复杂工业场景完成控制系统设计与部署", "items": [],
            }],
        }]

        html = render_resume_to_html(data)
        self.assertIn('<span class="project-inline-label">项目简介：</span>', html)
        self.assertNotIn('project-inline-label is-bold">项目简介', html)
        self.assertIn("text-align: justify", html)

        document = Document(BytesIO(generate_docx(data)))
        paragraph = next(p for p in document.paragraphs if "项目简介" in p.text)
        label_run = next(run for run in paragraph.runs if "项目简介" in run.text)
        self.assertFalse(label_run.bold)
        self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_skill_list_style_replaces_imported_markers_and_marker_weight_follows_content(self):
        data = resume_with_two_jobs()
        data["others"]["skills"] = ["1. Python 与 FastAPI", "**沟通协作**"]
        layout = default_layout_config()
        layout["skills"]["listStyle"] = "numbered"

        html = render_resume_to_html(data, layout_config=layout)

        self.assertIn('<ul class="list-items module-list list-style-numbered">', html)
        self.assertIn('<li class="list-item skill-list-item">Python 与 FastAPI</li>', html)
        self.assertIn('<li class="list-item skill-list-item marker-bold"><strong>沟通协作</strong></li>', html)
        self.assertNotIn('native-marker', html)
        self.assertIn('.list-item::before', html)
        self.assertIn('font-weight: 400', html)
        self.assertIn('.list-item.marker-bold::before', html)

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        numbered = next(paragraph for paragraph in document.paragraphs if "Python" in paragraph.text)
        bold_numbered = next(paragraph for paragraph in document.paragraphs if "沟通协作" in paragraph.text)
        self.assertTrue(numbered.text.startswith("\t(1)\tPython"))
        self.assertTrue(bold_numbered.text.startswith("\t(2)\t沟通协作"))
        self.assertFalse(next(run for run in numbered.runs if run.text == "(1)").bold)
        self.assertTrue(next(run for run in bold_numbered.runs if run.text == "(2)").bold)
        self.assertGreater(numbered.paragraph_format.left_indent.mm, 0)
        self.assertAlmostEqual(
            numbered.paragraph_format.first_line_indent.mm,
            -numbered.paragraph_format.left_indent.mm,
            places=1,
        )
        self.assertAlmostEqual(numbered.paragraph_format.left_indent.mm, 6.1, places=1)

    def test_word_numbered_semantic_blocks_use_real_hanging_indent(self):
        data = resume_with_two_jobs()
        data["project_experience"] = [{
            "project_name": "Robot control",
            "date_range": ["2025.01", "2025.06"],
            "content_blocks": [{
                "type": "numbered_list",
                "label": "Responsibilities",
                "items": ["A deliberately long responsibility that wraps onto another visual line in Word."],
            }],
        }]

        document = Document(BytesIO(generate_docx(data)))
        numbered = next(paragraph for paragraph in document.paragraphs if "(1)" in paragraph.text)
        label = next(paragraph for paragraph in document.paragraphs if "Responsibilities" in paragraph.text)

        self.assertGreater(numbered.paragraph_format.left_indent.mm, 0)
        self.assertLess(numbered.paragraph_format.first_line_indent.mm, 0)
        self.assertGreater(
            numbered.paragraph_format.left_indent.mm,
            label.paragraph_format.left_indent.mm,
        )
        self.assertEqual(numbered.text.count("\t"), 2)
        self.assertNotIn("(1) ", numbered.text)

    def test_word_disables_punctuation_overflow_on_measured_paragraphs(self):
        data = resume_with_two_jobs()
        data["project_experience"] = [{
            "project_name": "Punctuation boundary",
            "content_blocks": [{
                "type": "paragraph",
                "label": "项目简介",
                "text": "介绍，介绍。",
            }],
        }]

        with ZipFile(BytesIO(generate_docx(data))) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")

        self.assertIn('<w:overflowPunct w:val="0"', document_xml)

    def test_word_contact_details_use_dark_gray(self):
        data = resume_with_two_jobs()
        data["basics"].update({"birth_date": "2002.06", "phone": "17622312238", "email": "user@example.com"})

        document = Document(BytesIO(generate_docx(data)))
        contact_run = next(
            run
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if "user@example.com" in run.text
        )

        self.assertEqual(str(contact_run.font.color.rgb), "333333")

    def test_word_uses_the_same_body_type_scale_as_pdf_preview(self):
        data = resume_with_two_jobs()
        data["others"]["skills"] = ["Python 与 FastAPI"]
        layout = default_layout_config()
        layout["global"]["fontSize"] = 10.5

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        body_paragraph = next(paragraph for paragraph in document.paragraphs if "Python" in paragraph.text)

        self.assertAlmostEqual(body_paragraph.runs[-1].font.size.pt, 10.5, places=2)

    def test_semantic_font_sizes_and_inline_bold_match_pdf_and_word(self):
        data = resume_with_two_jobs()
        data["basics"]["name"] = "测试用户"
        data["work_experience"] = [{
            "company_name": "示例科技",
            "job_title": "后端工程师",
            "job_type": "实习",
            "date_range": ["2025.01", "2025.06"],
            "details": [],
            "content_blocks": [{
                "type": "paragraph",
                "label": "成果",
                "label_bold": True,
                "text": "将接口**延迟降低35%**并稳定运行",
                "items": [],
            }],
        }]
        layout = default_layout_config()
        layout["global"]["fontSize"] = 10
        layout["typography"]["fontSizes"] = {
            "name": 18.5,
            "sectionTitle": 13.5,
            "entryTitle": 12.5,
            "meta": 10.5,
            "body": 10,
            "label": 11.5,
        }

        html = render_resume_to_html(data, layout_config=layout)
        for declaration in (
            "--name-font-size: 18.5pt;",
            "--section-title-font-size: 13.5pt;",
            "--entry-title-font-size: 12.5pt;",
            "--meta-font-size: 10.5pt;",
            "--body-font-size: 10pt;",
            "--label-font-size: 11.5pt;",
        ):
            self.assertIn(declaration, html)
        self.assertIn("将接口<strong>延迟降低35%</strong>并稳定运行", html)

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        all_paragraphs = [paragraph for paragraph in document.paragraphs]
        all_paragraphs.extend(
            paragraph
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        name_run = next(run for paragraph in all_paragraphs for run in paragraph.runs if run.text == "测试用户")
        company_run = next(run for paragraph in all_paragraphs for run in paragraph.runs if run.text == "示例科技")
        date_run = next(run for paragraph in all_paragraphs for run in paragraph.runs if "2025.01" in run.text)
        # Unknown generic labels are flattened into the first content line so
        # the three-role contract cannot create an extra semantic block.
        label_run = next(run for paragraph in all_paragraphs for run in paragraph.runs if run.text == "成果")
        bold_body_run = next(run for paragraph in all_paragraphs for run in paragraph.runs if run.text == "延迟降低35%")
        section_run = next(
            paragraph.runs[0]
            for paragraph in document.paragraphs
            if paragraph.runs and paragraph._p.xpath("./w:pPr/w:pBdr")
        )
        self.assertAlmostEqual(name_run.font.size.pt, 18.5, places=1)
        self.assertAlmostEqual(section_run.font.size.pt, 13.5, places=1)
        self.assertAlmostEqual(company_run.font.size.pt, 12.5, places=1)
        self.assertAlmostEqual(date_run.font.size.pt, 11.5, places=1)
        self.assertAlmostEqual(label_run.font.size.pt, 10, places=1)
        self.assertAlmostEqual(bold_body_run.font.size.pt, 10, places=1)
        self.assertTrue(bold_body_run.bold)

    def test_pdf_inline_formatting_escapes_html_before_rendering_bold(self):
        data = resume_with_two_jobs()
        data["basics"]["name"] = '<img src=x onerror="boom"> **张三**'
        html = render_resume_to_html(data)
        self.assertIn('&lt;img src=x onerror=&quot;boom&quot;&gt; <strong>张三</strong>', html)
        self.assertNotIn('<img src=x onerror="boom">', html)

    def test_word_left_aligned_header_uses_full_row_when_photo_is_absent(self):
        data = resume_with_two_jobs()
        data["basics"].update({
            "birth_date": "2002.06",
            "phone": "17622312238",
            "email": "2776553477@qq.com",
        })
        layout = default_layout_config()
        layout["basics"]["preset"] = "left-aligned"

        with ZipFile(BytesIO(generate_docx(data, layout_config=layout))) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")

        self.assertIn('<w:tblGrid><w:gridCol w:w="10885"', document_xml)
        self.assertNotIn("<w:docGrid", document_xml)

    def test_one_line_education_layout_allows_long_columns_to_wrap_in_bounds(self):
        html = render_resume_to_html(resume_with_two_jobs())

        self.assertIn("overflow-wrap: break-word", html)
        self.assertIn("--letter-spacing: 0pt", html)
        self.assertIn("letter-spacing: var(--letter-spacing)", html)
        self.assertIn("font-kerning: none", html)
        self.assertIn("font-variant-ligatures: none", html)
        self.assertIn("font-synthesis: none", html)
        self.assertIn("--education-side-column: 42mm", html)
        self.assertIn("grid-template-columns: var(--education-side-column) minmax(0, 1fr) var(--education-side-column)", html)
        self.assertIn(".education-item .education-middle-column", html)
        self.assertIn("text-align: left", html)
        self.assertIn("margin-right: 0", html)
        self.assertIn("position: static", html)
        self.assertNotIn("right: 6mm", html)

    def test_pdf_bullet_marker_is_centered_without_moving_the_text_column(self):
        html = render_resume_to_html(resume_with_two_jobs())

        self.assertIn("padding-left: var(--list-text-indent)", html)
        self.assertIn("width: calc(var(--list-text-indent) - var(--list-marker-gap))", html)
        self.assertIn("text-align: center", html)

    def test_word_one_line_education_uses_full_printable_width_without_right_padding(self):
        data = resume_with_two_jobs()
        data["education"] = [{
            "school_name": "暨南大学",
            "degree": "硕士",
            "major": "电子信息",
            "date_range": ["2024.09", "2027.06"],
        }]
        layout = default_layout_config()
        layout["education"]["preset"] = "compact"

        with ZipFile(BytesIO(generate_docx(data, layout_config=layout))) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        education_table = next(
            table for table in document.tables
            if any("暨南大学" in cell.text for cell in table.rows[0].cells)
        )

        self.assertIn('<w:tblLayout w:type="fixed"', document_xml)
        self.assertIn('<w:tblW w:type="dxa" w:w="10885"', document_xml)
        self.assertNotIn('<w:end w:w="80" w:type="dxa"', document_xml)
        self.assertEqual(len(education_table.columns), 3)
        self.assertEqual(education_table.cell(0, 1).text, "硕士 · 电子信息")
        self.assertAlmostEqual(
            education_table.columns[0].width.mm,
            education_table.columns[2].width.mm,
            places=1,
        )
        self.assertAlmostEqual(
            sum(column.width.mm for column in education_table.columns),
            192.0,
            places=1,
        )
        self.assertEqual(education_table.cell(0, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(education_table.cell(0, 2).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_compact_education_keeps_unlabeled_metrics_in_left_aligned_middle_frame(self):
        data = resume_with_two_jobs()
        data["education"] = [{
            "school_name": "中山大学",
            "degree": "硕士",
            "major": "电子信息",
            "gpa": "4.0",
            "gpa_scale": "5.0",
            "ranking": "前5%",
            "date_range": ["2024.09", "2027.06"],
        }]
        layout = default_layout_config()

        html = render_resume_to_html(data, layout_config=layout)
        self.assertIn('component-degree">硕士</span>', html)
        self.assertIn('component-major">电子信息</span>', html)
        self.assertIn('component-metrics">4.0/5.0 (前5%)</span>', html)
        self.assertIn("--education-compact-side-column:", html)
        self.assertIn("--education-middle-column:", html)

        content = generate_docx(data, layout_config=layout)
        document = Document(BytesIO(content))
        education_table = next(
            table for table in document.tables
            if any("中山大学" in cell.text for cell in table.rows[0].cells)
        )
        self.assertEqual(len(education_table.columns), 3)
        self.assertEqual(
            education_table.cell(0, 1).text,
            "硕士 · 电子信息 · 4.0/5.0 (前5%)",
        )
        self.assertEqual(education_table.cell(0, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertAlmostEqual(
            education_table.columns[0].width.mm,
            education_table.columns[2].width.mm,
            places=1,
        )

        with ZipFile(BytesIO(content)) as archive:
            settings_xml = archive.read("word/settings.xml").decode("utf-8")
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('<w:characterSpacingControl w:val="doNotCompress"', settings_xml)
        self.assertIn('w:eastAsia="zh-CN"', settings_xml)
        self.assertIn('<w:overflowPunct w:val="0"', document_xml)

    def test_mixed_text_keeps_source_spaces_and_section_divider_gap_matches(self):
        data = resume_with_two_jobs()
        mixed_text = "混排正文与 ASCII token 保留普通空格和自然换行"
        data["work_experience"][0]["details"] = [mixed_text]
        layout = default_layout_config()

        html = render_resume_to_html(data, layout_config=layout)
        self.assertIn(mixed_text, html)
        self.assertIn("--section-title-border-gap: 1.1pt", html)
        self.assertIn("padding-bottom: var(--section-title-border-gap)", html)

        document = Document(BytesIO(generate_docx(data, layout_config=layout)))
        body = next(paragraph for paragraph in document.paragraphs if "ASCII token" in paragraph.text)
        self.assertIn(mixed_text, body.text)
        self.assertNotIn("\u00a0", body.text)
        self.assertNotIn("\u2060", body.text)
        heading = next(
            paragraph for paragraph in document.paragraphs
            if paragraph.runs and paragraph._p.xpath("./w:pPr/w:pBdr")
        )
        border = heading._p.xpath("./w:pPr/w:pBdr/w:bottom")[0]
        self.assertEqual(border.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space"), "1")


if __name__ == "__main__":
    unittest.main()
