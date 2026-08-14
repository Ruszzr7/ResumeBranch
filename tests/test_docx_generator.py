import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from backend.docx_generator import generate_docx
from backend.llm_providers import PROVIDERS, public_registry, role_temperature, temperature_supported, validate_profile
from backend import resume_agent


class DocxGeneratorTests(unittest.TestCase):
    def test_mixed_punctuation_keeps_codepoints_and_script_font_mapping(self):
        punctuation_sample = "中文，句号。分号；冒号：括号（）/ ASCII,.;:!?() RL部署经验"
        document = Document(BytesIO(generate_docx({
            "basics": {"name": "标点测试"},
            "self_evaluation": [punctuation_sample],
        })))
        paragraph = next(item for item in document.paragraphs if punctuation_sample in item.text)
        self.assertEqual(paragraph.text, punctuation_sample)
        for run in paragraph.runs:
            fonts = run._element.rPr.rFonts
            self.assertEqual(fonts.get(qn("w:ascii")), "Arial")
            self.assertEqual(fonts.get(qn("w:hAnsi")), "Arial")
            self.assertEqual(fonts.get(qn("w:eastAsia")), "Microsoft YaHei")

    def test_generates_editable_resume_with_academic_metrics(self):
        content = generate_docx({
            "basics": {"name": "测试用户", "target_position": "产品经理"},
            "education": [{
                "school_name": "示例大学",
                "degree": "本科",
                "major": "计算机科学",
                "gpa": "3.72",
                "gpa_scale": "4.0",
                "ranking": "前 10%",
                "date_range": ["2022", "2025"],
            }],
            "work_experience": [],
            "project_experience": [],
            "others": {},
            "self_evaluation": [],
        })

        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        combined = text + "\n" + table_text
        self.assertIn("测试用户", combined)
        self.assertIn("示例大学", combined)
        self.assertIn("3.72/4.0 (前 10%)", combined)
        self.assertNotIn("GPA：3.72/4.0", combined)
        self.assertAlmostEqual(document.sections[0].page_width.mm, 210.0, places=1)
        normal_fonts = document.styles["Normal"]._element.rPr.rFonts
        self.assertEqual(normal_fonts.get(qn("w:ascii")), "Arial")
        self.assertEqual(normal_fonts.get(qn("w:hAnsi")), "Arial")
        self.assertEqual(normal_fonts.get(qn("w:eastAsia")), "Microsoft YaHei")

        name_run = next(
            run
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if "测试用户" in run.text
        )
        self.assertEqual(name_run._element.rPr.rFonts.get(qn("w:ascii")), "Arial")
        self.assertEqual(name_run._element.rPr.rFonts.get(qn("w:eastAsia")), "Microsoft YaHei")
        self.assertEqual(name_run._element.rPr.xpath("./w:spacing")[0].get(qn("w:val")), "0")
        self.assertEqual(name_run._element.rPr.xpath("./w:kern")[0].get(qn("w:val")), "0")
        self.assertAlmostEqual(name_run.font.size.pt, 14, places=1)
        self.assertAlmostEqual(document.styles["Normal"].font.size.pt, 9, places=1)

        bold_table_sizes = [
            run.font.size.pt
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if run.bold and run.font.size is not None
        ]
        self.assertTrue(any(abs(size - 10) < 0.1 for size in bold_table_sizes))
        section_heading = next(
            paragraph.runs[0]
            for paragraph in document.paragraphs
            if paragraph.runs
            and paragraph._p.xpath("./w:pPr/w:pBdr")
        )
        self.assertAlmostEqual(section_heading.font.size.pt, 11, places=1)

        # Without a photo, the centered header uses the full printable width so
        # contact details do not wrap inside the old narrow middle column.
        self.assertEqual(len(document.tables[0]._tbl.tr_lst[0].tc_lst), 1)

        education_table = next(
            table for table in document.tables
            if any("示例大学" in cell.text for row in table.rows for cell in row.cells)
        )
        education_header = " ".join(cell.text for row in education_table.rows for cell in row.cells)
        self.assertIn("本科", education_header)
        self.assertIn("本科 · 计算机科学", education_header)
        education_widths = [cell.width.mm for cell in education_table.rows[0].cells]
        self.assertAlmostEqual(education_widths[0], education_widths[2], places=1)
        self.assertLess(education_widths[1], education_widths[0])


class ProviderRulesTests(unittest.TestCase):
    def test_public_provider_registry_does_not_publish_stale_models_or_urls(self):
        for definition in PROVIDERS.values():
            self.assertNotIn("models", definition)
            self.assertNotIn("default_model", definition)
        for provider in public_registry():
            self.assertNotIn("models", provider)
            self.assertNotIn("default_model", provider)
            self.assertNotIn("base_url", provider)

    def test_model_specific_temperature_rules(self):
        self.assertTrue(temperature_supported("kimi_coding", "k3"))
        self.assertFalse(temperature_supported("openai", "gpt-5.2"))
        self.assertTrue(temperature_supported("openai", "gpt-4.1"))
        self.assertFalse(temperature_supported("deepseek", "deepseek-reasoner"))
        self.assertEqual(role_temperature("openai_chat", "k3-256k", 0.0), 1.0)
        self.assertEqual(role_temperature("openai_chat", "kimi-for-coding", 0.1), 1.0)

    def test_validation_no_longer_depends_on_vendor_temperature_ranges(self):
        validate_profile("openai_chat", "glm-5.2", "https://open.bigmodel.cn/api/paas/v4/", 0.0)

    def test_role_temperature_is_not_overridden_by_saved_profile(self):
        with (
            patch.object(resume_agent, "LLM_TEMPERATURE", 1.0),
            patch.object(resume_agent, "create_llm_for_config", return_value=object()) as factory,
        ):
            resume_agent.create_llm(temperature=0.0)

        self.assertEqual(factory.call_args.kwargs["temperature"], 0.0)


if __name__ == "__main__":
    unittest.main()
